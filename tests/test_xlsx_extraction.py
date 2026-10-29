from datetime import datetime

from openpyxl import Workbook

from rag.chunk_strategy import ChunkStrategyRouter
from utils.file_handler import xlsx_loader


def test_xlsx_loader_preserves_headers_values_and_locations(tmp_path):
    filepath = tmp_path / "products.xlsx"
    workbook = Workbook()
    products = workbook.active
    products.title = "产品清单"
    products.append(["产品型号", "价格", "是否在售", "更新时间", "库存公式"])
    products.append(["KB-R2", 2999, True, datetime(2026, 9, 19, 10, 30), "=1+2"])
    products.append([None, None, None, None, None])
    products.append(["KB-R3", 3999, False, datetime(2026, 9, 20), "=2+3"])

    service = workbook.create_sheet("售后政策")
    service.append(["故障等级", "响应时间"])
    service.append(["P0", "10分钟"])
    workbook.create_sheet("空表")
    workbook.save(filepath)
    workbook.close()

    documents = xlsx_loader(str(filepath))

    assert len(documents) == 5
    assert {doc.metadata["sheet_name"] for doc in documents} == {"产品清单", "售后政策"}

    kb_r2 = next(doc for doc in documents if "KB-R2" in doc.page_content)
    assert "产品型号: KB-R2" in kb_r2.page_content
    assert "价格: 2999" in kb_r2.page_content
    assert "是否在售: 是" in kb_r2.page_content
    assert "更新时间: 2026-09-19 10:30:00" in kb_r2.page_content
    assert "库存公式: =1+2" in kb_r2.page_content
    assert kb_r2.metadata == {
        "source": str(filepath),
        "block_type": "spreadsheet_row",
        "sheet_name": "产品清单",
        "sheet_index": 0,
        "row_number": 2,
        "column_count": 5,
    }

    p0 = next(doc for doc in documents if "P0" in doc.page_content)
    assert "工作表：售后政策" in p0.page_content
    assert "故障等级: P0" in p0.page_content
    assert "响应时间: 10分钟" in p0.page_content
    assert p0.metadata["row_number"] == 2

    chunks = ChunkStrategyRouter(chunk_size=80, chunk_overlap=10).split(
        documents,
        source_type="xlsx",
    )
    assert chunks
    assert all(chunk.metadata.get("sheet_name") for chunk in chunks)
    assert any("KB-R2" in chunk.page_content for chunk in chunks)


def test_xlsx_loader_handles_title_row_and_empty_workbook(tmp_path):
    titled_path = tmp_path / "titled.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "统计"
    sheet.append(["2026年度统计"])
    sheet.append(["部门", "人数"])
    sheet.append(["研发部", 35])
    workbook.save(titled_path)
    workbook.close()

    documents = xlsx_loader(str(titled_path))

    title = next(doc for doc in documents if "2026年度统计" in doc.page_content)
    assert "A: 2026年度统计" in title.page_content
    employee_count = next(doc for doc in documents if "研发部" in doc.page_content)
    assert "部门: 研发部" in employee_count.page_content
    assert "人数: 35" in employee_count.page_content

    empty_path = tmp_path / "empty.xlsx"
    empty_workbook = Workbook()
    empty_workbook.save(empty_path)
    empty_workbook.close()

    assert xlsx_loader(str(empty_path)) == []
