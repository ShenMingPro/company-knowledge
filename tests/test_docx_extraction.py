"""
DOCX 结构化文本提取测试
测试标题、段落、表格的提取和顺序保持
"""
import os
import sys
import tempfile
import shutil

# 添加项目根目录到路径
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document as DocxDocument
from docx.shared import Pt
from utils.file_handler import docx_loader


def create_test_docx_with_structure(filepath: str):
    """创建包含标题、段落、表格的测试 DOCX"""
    doc = DocxDocument()

    # 添加标题 1
    h1 = doc.add_heading('第一章 产品介绍', level=1)

    # 添加段落
    doc.add_paragraph('本产品适用于家庭场景，支持多种工作模式。')

    # 添加标题 2
    doc.add_heading('1.1 产品特点', level=2)

    # 添加段落
    doc.add_paragraph('支持自动回充、智能路径规划等功能。')

    # 添加表格
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = '字段'
    table.cell(0, 1).text = '内容'
    table.cell(1, 0).text = '型号'
    table.cell(1, 1).text = 'KB-R2'
    table.cell(2, 0).text = '故障码'
    table.cell(2, 1).text = 'E03'

    # 添加标题 2
    doc.add_heading('1.2 技术参数', level=2)

    # 添加段落
    doc.add_paragraph('工作电压：220V，功率：1500W。')

    doc.save(filepath)


def create_empty_docx(filepath: str):
    """创建空白 DOCX"""
    doc = DocxDocument()
    doc.save(filepath)


def create_docx_with_empty_paragraphs(filepath: str):
    """创建包含空段落的 DOCX"""
    doc = DocxDocument()
    doc.add_paragraph('')
    doc.add_paragraph('   ')
    doc.add_paragraph('有效内容')
    doc.add_paragraph('')
    doc.save(filepath)


def create_docx_with_text_headings(filepath: str):
    """创建使用文本规则识别标题的 DOCX（无样式）"""
    doc = DocxDocument()

    # Markdown 风格
    doc.add_paragraph('## 使用说明')
    doc.add_paragraph('这是说明内容。')

    # 数字编号
    doc.add_paragraph('1. 产品介绍')
    doc.add_paragraph('产品描述内容。')

    # 中文章节
    doc.add_paragraph('一、安装步骤')
    doc.add_paragraph('安装说明。')

    doc.save(filepath)


def create_docx_with_normal_text(filepath: str):
    """创建只有普通正文的 DOCX（避免误判）"""
    doc = DocxDocument()

    doc.add_paragraph('这是一段普通的正文，比较短。')
    doc.add_paragraph('这是另一段正文，包含型号 KB-R2 和故障码 E03。')
    doc.add_paragraph('这段正文稍微长一点，用来测试不会被误判为标题。')

    doc.save(filepath)


def test_heading_paragraph_table_extraction():
    """测试标题、段落、表格的提取"""
    print("\n=== 测试标题、段落、表格的提取 ===")

    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = os.path.join(tmpdir, 'test.docx')
        create_test_docx_with_structure(filepath)

        docs = docx_loader(filepath)

        print(f"提取到 {len(docs)} 个内容块")
        assert len(docs) > 0, "应该提取到内容块"

        # 检查内容块类型
        block_types = [doc.metadata['block_type'] for doc in docs]
        assert 'heading' in block_types, "应该包含标题块"
        assert 'paragraph' in block_types, "应该包含段落块"
        assert 'table' in block_types, "应该包含表格块"

        # 检查标题块
        headings = [doc for doc in docs if doc.metadata['block_type'] == 'heading']
        assert len(headings) >= 3, "应该有至少 3 个标题"

        for h in headings:
            assert 'heading_level' in h.metadata, "标题应该有 heading_level"
            assert h.metadata['is_heading'] is True, "标题的 is_heading 应为 True"
            assert 'section_path' in h.metadata, "标题应该有 section_path"
            print(f"  标题 [{h.metadata['heading_level']}]: {h.page_content[:30]}")

        # 检查段落块
        paragraphs = [doc for doc in docs if doc.metadata['block_type'] == 'paragraph']
        assert len(paragraphs) >= 3, "应该有至少 3 个段落"

        for p in paragraphs:
            assert 'paragraph_index' in p.metadata, "段落应该有 paragraph_index"
            assert p.metadata['is_heading'] is False, "段落的 is_heading 应为 False"
            assert 'section_path' in p.metadata, "段落应该有 section_path"

        # 检查表格块
        tables = [doc for doc in docs if doc.metadata['block_type'] == 'table']
        assert len(tables) >= 1, "应该有至少 1 个表格"

        for t in tables:
            assert 'table_index' in t.metadata, "表格应该有 table_index"
            assert 'row_count' in t.metadata, "表格应该有 row_count"
            assert 'column_count' in t.metadata, "表格应该有 column_count"
            assert 'section_path' in t.metadata, "表格应该有 section_path"
            print(f"  表格 [{t.metadata['row_count']}x{t.metadata['column_count']}]")

        print("✓ 标题、段落、表格提取通过")


def test_content_order_preservation():
    """测试内容顺序保持"""
    print("\n=== 测试内容顺序保持 ===")

    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = os.path.join(tmpdir, 'test.docx')
        create_test_docx_with_structure(filepath)

        docs = docx_loader(filepath)

        # 检查顺序：应该是 heading -> paragraph -> heading -> paragraph -> table -> heading -> paragraph
        expected_sequence = [
            'heading',    # 第一章 产品介绍
            'paragraph',  # 本产品适用于...
            'heading',    # 1.1 产品特点
            'paragraph',  # 支持自动回充...
            'table',      # 型号表格
            'heading',    # 1.2 技术参数
            'paragraph'   # 工作电压...
        ]

        actual_sequence = [doc.metadata['block_type'] for doc in docs]

        print(f"期望顺序: {expected_sequence}")
        print(f"实际顺序: {actual_sequence}")

        assert actual_sequence == expected_sequence, f"内容顺序不匹配"

        print("✓ 内容顺序保持通过")


def test_heading_style_recognition():
    """测试 Word 样式标题识别"""
    print("\n=== 测试 Word 样式标题识别 ===")

    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = os.path.join(tmpdir, 'test.docx')
        create_test_docx_with_structure(filepath)

        docs = docx_loader(filepath)
        headings = [doc for doc in docs if doc.metadata['block_type'] == 'heading']

        # 检查标题级别
        levels = [h.metadata['heading_level'] for h in headings]
        assert 1 in levels, "应该识别到 1 级标题"
        assert 2 in levels, "应该识别到 2 级标题"

        # 检查样式名称
        for h in headings:
            style_name = h.metadata.get('style_name', '')
            print(f"  {h.page_content[:20]} - 级别: {h.metadata['heading_level']}, 样式: {style_name}")

        print("✓ Word 样式标题识别通过")


def test_text_heading_recognition():
    """测试文本规则标题识别"""
    print("\n=== 测试文本规则标题识别 ===")

    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = os.path.join(tmpdir, 'test.docx')
        create_docx_with_text_headings(filepath)

        docs = docx_loader(filepath)

        # 应该识别出文本风格的标题
        headings = [doc for doc in docs if doc.metadata['block_type'] == 'heading']

        print(f"识别到 {len(headings)} 个文本标题")
        for h in headings:
            print(f"  {h.page_content} - 样式: {h.metadata.get('style_name', '')}")

        # 至少应该识别一些标题
        assert len(headings) > 0, "应该识别出文本风格的标题"

        print("✓ 文本规则标题识别通过")


def test_normal_text_not_misidentified():
    """测试普通正文不会被误判为标题"""
    print("\n=== 测试普通正文不会被误判 ===")

    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = os.path.join(tmpdir, 'test.docx')
        create_docx_with_normal_text(filepath)

        docs = docx_loader(filepath)

        # 所有内容应该被识别为段落
        headings = [doc for doc in docs if doc.metadata['block_type'] == 'heading']
        paragraphs = [doc for doc in docs if doc.metadata['block_type'] == 'paragraph']

        print(f"识别为标题: {len(headings)}, 识别为段落: {len(paragraphs)}")

        assert len(headings) == 0, "普通正文不应该被误判为标题"
        assert len(paragraphs) >= 3, "应该有 3 个段落"

        print("✓ 普通正文不会被误判通过")


def test_section_path_inheritance():
    """测试章节路径继承"""
    print("\n=== 测试章节路径继承 ===")

    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = os.path.join(tmpdir, 'test.docx')
        create_test_docx_with_structure(filepath)

        docs = docx_loader(filepath)

        for doc in docs:
            block_type = doc.metadata['block_type']
            section_path = doc.metadata.get('section_path', '')
            content_preview = doc.page_content[:30].replace('\n', ' ')
            print(f"  [{block_type:10}] {content_preview:40} -> {section_path}")

        # 检查第一个段落的 section_path 应该包含第一章
        first_para = [doc for doc in docs if doc.metadata['block_type'] == 'paragraph'][0]
        assert '第一章 产品介绍' in first_para.metadata['section_path'], \
            "第一个段落应该属于第一章"

        # 检查 1.1 节下的段落
        second_para = [doc for doc in docs if doc.metadata['block_type'] == 'paragraph'][1]
        section_path = second_para.metadata['section_path']
        assert '第一章 产品介绍' in section_path, "应该包含第一章"
        assert '1.1 产品特点' in section_path, "应该包含 1.1 节"

        print("✓ 章节路径继承通过")


def test_table_content_preservation():
    """测试表格内容保留"""
    print("\n=== 测试表格内容保留 ===")

    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = os.path.join(tmpdir, 'test.docx')
        create_test_docx_with_structure(filepath)

        docs = docx_loader(filepath)
        tables = [doc for doc in docs if doc.metadata['block_type'] == 'table']

        assert len(tables) >= 1, "应该有至少 1 个表格"

        table_content = tables[0].page_content
        print(f"表格内容:\n{table_content}")

        # 检查表格内容
        assert 'KB-R2' in table_content, "表格应该包含型号 KB-R2"
        assert 'E03' in table_content, "表格应该包含故障码 E03"
        assert '型号' in table_content, "表格应该包含字段名"

        # 检查行列关系（使用 | 分隔）
        assert '|' in table_content, "表格应该使用 | 分隔单元格"

        # 检查行数列数
        assert tables[0].metadata['row_count'] == 3, "表格应该有 3 行"
        assert tables[0].metadata['column_count'] == 2, "表格应该有 2 列"

        print("✓ 表格内容保留通过")


def test_model_number_preservation():
    """测试型号和故障码保持完整"""
    print("\n=== 测试型号和故障码保持完整 ===")

    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = os.path.join(tmpdir, 'test.docx')
        create_test_docx_with_structure(filepath)

        docs = docx_loader(filepath)

        # 在表格中查找型号
        tables = [doc for doc in docs if doc.metadata['block_type'] == 'table']
        table_content = tables[0].page_content

        assert 'KB-R2' in table_content, "型号 KB-R2 应该保持完整"
        assert 'E03' in table_content, "故障码 E03 应该保持完整"

        # 检查没有被破坏
        assert 'KB - R2' not in table_content, "型号不应该被多余空格破坏"
        assert 'KB-R 2' not in table_content, "型号不应该被多余空格破坏"

        print("✓ 型号和故障码保持完整通过")


def test_empty_paragraph_skipping():
    """测试空段落不生成内容块"""
    print("\n=== 测试空段落不生成内容块 ===")

    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = os.path.join(tmpdir, 'test.docx')
        create_docx_with_empty_paragraphs(filepath)

        docs = docx_loader(filepath)

        # 应该只有 1 个有效段落
        paragraphs = [doc for doc in docs if doc.metadata['block_type'] == 'paragraph']
        assert len(paragraphs) == 1, "空段落不应该生成内容块"
        assert paragraphs[0].page_content == '有效内容', "应该只保留有效内容"

        print("✓ 空段落不生成内容块通过")


def test_empty_docx():
    """测试空白 DOCX"""
    print("\n=== 测试空白 DOCX ===")

    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = os.path.join(tmpdir, 'empty.docx')
        create_empty_docx(filepath)

        docs = docx_loader(filepath)

        print(f"空白 DOCX 提取到 {len(docs)} 个内容块")
        assert len(docs) == 0, "空白 DOCX 应该返回空列表"

        print("✓ 空白 DOCX 通过")


def test_metadata_compatibility():
    """测试 metadata 与现有系统兼容"""
    print("\n=== 测试 metadata 与现有系统兼容 ===")

    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = os.path.join(tmpdir, 'test.docx')
        create_test_docx_with_structure(filepath)

        docs = docx_loader(filepath)

        # 检查所有必需字段都存在
        for doc in docs:
            assert 'block_type' in doc.metadata, "应该有 block_type"
            assert 'source' in doc.metadata, "应该有 source"
            assert 'section_path' in doc.metadata, "应该有 section_path"

            # 检查字段值类型
            assert isinstance(doc.metadata['block_type'], str), "block_type 应该是字符串"
            assert isinstance(doc.metadata['section_path'], str), "section_path 应该是字符串"

            # 标题块特有字段
            if doc.metadata['block_type'] == 'heading':
                assert 'heading_level' in doc.metadata, "标题应该有 heading_level"
                assert isinstance(doc.metadata['heading_level'], int), "heading_level 应该是整数"
                assert doc.metadata['is_heading'] is True, "标题的 is_heading 应为 True"

            # 段落块特有字段
            if doc.metadata['block_type'] == 'paragraph':
                assert 'paragraph_index' in doc.metadata, "段落应该有 paragraph_index"
                assert isinstance(doc.metadata['paragraph_index'], int), "paragraph_index 应该是整数"
                assert doc.metadata['is_heading'] is False, "段落的 is_heading 应为 False"

            # 表格块特有字段
            if doc.metadata['block_type'] == 'table':
                assert 'table_index' in doc.metadata, "表格应该有 table_index"
                assert 'row_count' in doc.metadata, "表格应该有 row_count"
                assert 'column_count' in doc.metadata, "表格应该有 column_count"

        print("✓ metadata 与现有系统兼容通过")


def run_all_tests():
    """运行所有测试"""
    print("=" * 60)
    print("DOCX 结构化文本提取测试")
    print("=" * 60)

    tests = [
        test_heading_paragraph_table_extraction,
        test_content_order_preservation,
        test_heading_style_recognition,
        test_text_heading_recognition,
        test_normal_text_not_misidentified,
        test_section_path_inheritance,
        test_table_content_preservation,
        test_model_number_preservation,
        test_empty_paragraph_skipping,
        test_empty_docx,
        test_metadata_compatibility,
    ]

    passed = 0
    failed = 0

    for test in tests:
        try:
            test()
            passed += 1
        except AssertionError as e:
            print(f"✗ {test.__name__} 失败: {e}")
            failed += 1
        except Exception as e:
            print(f"✗ {test.__name__} 出错: {e}")
            failed += 1

    print("\n" + "=" * 60)
    print(f"测试完成: {passed} 通过, {failed} 失败")
    print("=" * 60)

    return failed == 0


if __name__ == '__main__':
    success = run_all_tests()
    sys.exit(0 if success else 1)
