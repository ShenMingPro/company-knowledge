"""
DOCX 结构化提取与现有系统集成测试
验证新的 DOCX 提取能够与 vector_store 正常配合
"""
import os
import sys
import tempfile

# 添加项目根目录到路径
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document as DocxDocument
from utils.file_handler import docx_loader, pdf_loader, txt_loader
from langchain_text_splitters import RecursiveCharacterTextSplitter


def test_docx_returns_multiple_documents():
    """测试 DOCX 返回多个结构化 Document"""
    print("\n=== 测试 DOCX 返回多个结构化 Document ===")

    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = os.path.join(tmpdir, 'test.docx')

        # 创建测试文档
        doc = DocxDocument()
        doc.add_heading('产品介绍', level=1)
        doc.add_paragraph('这是产品描述。')
        doc.add_heading('技术参数', level=2)
        doc.add_paragraph('参数详情。')
        doc.save(filepath)

        # 使用新的 docx_loader
        docs = docx_loader(filepath)

        print(f"提取到 {len(docs)} 个 Document")
        assert len(docs) > 1, "应该返回多个 Document"

        # 每个 Document 都应该有 metadata
        for doc in docs:
            assert hasattr(doc, 'metadata'), "应该有 metadata"
            assert hasattr(doc, 'page_content'), "应该有 page_content"
            assert 'block_type' in doc.metadata, "metadata 应该有 block_type"

        print("✓ DOCX 返回多个结构化 Document 通过")


def test_docx_with_text_splitter():
    """测试 DOCX 提取结果能够被 RecursiveCharacterTextSplitter 处理"""
    print("\n=== 测试 DOCX 与 TextSplitter 配合 ===")

    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = os.path.join(tmpdir, 'test.docx')

        # 创建较长的文档
        doc = DocxDocument()
        doc.add_heading('第一章 产品介绍', level=1)
        doc.add_paragraph('这是第一段。' * 50)  # 较长的段落
        doc.add_heading('第二章 使用说明', level=1)
        doc.add_paragraph('这是第二段。' * 50)
        doc.save(filepath)

        # 使用 docx_loader
        docs = docx_loader(filepath)
        print(f"提取到 {len(docs)} 个原始 Document")

        # 使用 RecursiveCharacterTextSplitter（模拟 vector_store.py 的行为）
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=200,
            chunk_overlap=50,
            separators=["\n\n", "\n", " ", ""],
            length_function=len,
        )

        split_docs = splitter.split_documents(docs)
        print(f"分片后得到 {len(split_docs)} 个 Document")

        assert len(split_docs) > 0, "分片后应该有内容"

        # 检查 metadata 是否保留
        for split_doc in split_docs:
            assert 'block_type' in split_doc.metadata, "分片后应该保留 block_type"
            assert 'source' in split_doc.metadata, "分片后应该保留 source"

        print("✓ DOCX 与 TextSplitter 配合通过")


def test_metadata_not_overwritten():
    """测试 DOCX metadata 不会覆盖后续添加的字段"""
    print("\n=== 测试 metadata 不会被覆盖 ===")

    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = os.path.join(tmpdir, 'test.docx')

        doc = DocxDocument()
        doc.add_heading('测试标题', level=1)
        doc.add_paragraph('测试段落。')
        doc.save(filepath)

        docs = docx_loader(filepath)

        # 模拟 vector_store.py 添加的 metadata
        for idx, doc in enumerate(docs):
            doc.metadata.update({
                'space_id': 1,
                'document_id': 100,
                'document_version_id': 200,
                'source_name': 'test.docx',
                'chunk_index': idx,
                'is_active': True
            })

        # 验证所有字段都存在
        for doc in docs:
            # DOCX 原有字段
            assert 'block_type' in doc.metadata, "应该保留 block_type"
            assert 'section_path' in doc.metadata, "应该保留 section_path"

            # vector_store 添加的字段
            assert doc.metadata['space_id'] == 1, "space_id 应该正确"
            assert doc.metadata['document_id'] == 100, "document_id 应该正确"
            assert doc.metadata['document_version_id'] == 200, "version_id 应该正确"
            assert doc.metadata['chunk_index'] >= 0, "chunk_index 应该存在"
            assert doc.metadata['is_active'] is True, "is_active 应该正确"

        print("✓ metadata 不会被覆盖通过")


def test_pdf_loader_unchanged():
    """测试 PDF loader 保持不变"""
    print("\n=== 测试 PDF loader 保持不变 ===")

    # pdf_loader 应该仍然使用 PyPDFLoader
    # 这里只验证函数存在且可调用
    assert callable(pdf_loader), "pdf_loader 应该存在"

    print("✓ PDF loader 保持不变通过")


def test_txt_loader_unchanged():
    """测试 TXT loader 保持不变"""
    print("\n=== 测试 TXT loader 保持不变 ===")

    # txt_loader 应该仍然使用 TextLoader
    assert callable(txt_loader), "txt_loader 应该存在"

    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = os.path.join(tmpdir, 'test.txt')
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write('测试文本内容')

        docs = txt_loader(filepath)
        assert len(docs) > 0, "TXT loader 应该能正常工作"

    print("✓ TXT loader 保持不变通过")


def test_docx_metadata_fields():
    """测试 DOCX metadata 字段完整性"""
    print("\n=== 测试 DOCX metadata 字段完整性 ===")

    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = os.path.join(tmpdir, 'test.docx')

        doc = DocxDocument()
        doc.add_heading('标题', level=1)
        doc.add_paragraph('段落')

        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'A'
        table.cell(0, 1).text = 'B'

        doc.save(filepath)

        docs = docx_loader(filepath)

        # 标题块
        headings = [d for d in docs if d.metadata.get('block_type') == 'heading']
        assert len(headings) > 0, "应该有标题块"
        h = headings[0]
        assert 'heading_level' in h.metadata
        assert 'is_heading' in h.metadata
        assert h.metadata['is_heading'] is True
        assert 'style_name' in h.metadata
        assert 'section_path' in h.metadata
        assert 'source' in h.metadata

        # 段落块
        paragraphs = [d for d in docs if d.metadata.get('block_type') == 'paragraph']
        assert len(paragraphs) > 0, "应该有段落块"
        p = paragraphs[0]
        assert 'paragraph_index' in p.metadata
        assert 'is_heading' in p.metadata
        assert p.metadata['is_heading'] is False
        assert 'section_path' in p.metadata
        assert 'style_name' in p.metadata
        assert 'source' in p.metadata

        # 表格块
        tables = [d for d in docs if d.metadata.get('block_type') == 'table']
        assert len(tables) > 0, "应该有表格块"
        t = tables[0]
        assert 'table_index' in t.metadata
        assert 'row_count' in t.metadata
        assert 'column_count' in t.metadata
        assert 'section_path' in t.metadata
        assert 'source' in t.metadata

        print("✓ DOCX metadata 字段完整性通过")


def run_all_tests():
    """运行所有集成测试"""
    print("=" * 60)
    print("DOCX 结构化提取集成测试")
    print("=" * 60)

    tests = [
        test_docx_returns_multiple_documents,
        test_docx_with_text_splitter,
        test_metadata_not_overwritten,
        test_pdf_loader_unchanged,
        test_txt_loader_unchanged,
        test_docx_metadata_fields,
    ]

    passed = 0
    failed = 0

    for test in tests:
        try:
            test()
            passed += 1
        except AssertionError as e:
            print(f"✗ {test.__name__} 失败: {e}")
            failed += 1
        except Exception as e:
            print(f"✗ {test.__name__} 出错: {e}")
            import traceback
            traceback.print_exc()
            failed += 1

    print("\n" + "=" * 60)
    print(f"测试完成: {passed} 通过, {failed} 失败")
    print("=" * 60)

    return failed == 0


if __name__ == '__main__':
    success = run_all_tests()
    sys.exit(0 if success else 1)
