import re
from datetime import date, datetime, time
from typing import Optional
from utils.logger_handler import logger
from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def pdf_loader(filepath: str, passwd=None) -> list[Document]:
    return PyPDFLoader(filepath, passwd).load()


def txt_loader(filepath: str) -> list[Document]:
    return TextLoader(filepath, encoding="utf-8").load()


def xlsx_loader(filepath: str) -> list[Document]:
    """Parse an XLSX workbook into searchable row-level documents.

    The first table-like row in each sheet is treated as the header. Its values
    are repeated in subsequent row documents so that individual cells retain
    their column meaning after vector retrieval.
    """
    workbook = None
    try:
        workbook = load_workbook(filepath, read_only=True, data_only=False)
        documents = []

        for sheet_index, worksheet in enumerate(workbook.worksheets):
            headers = None

            for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                values = [_format_excel_value(value) for value in row]
                while values and not values[-1]:
                    values.pop()

                if not any(values):
                    continue

                if headers is None and _looks_like_excel_header(row):
                    headers = values
                    content = (
                        f"工作表：{worksheet.title}\n"
                        f"第{row_number}行（表头）：{' | '.join(value for value in values if value)}"
                    )
                    block_type = "spreadsheet_header"
                else:
                    fields = []
                    for column_index, value in enumerate(values, start=1):
                        if not value:
                            continue
                        header = (
                            headers[column_index - 1]
                            if headers and column_index <= len(headers) and headers[column_index - 1]
                            else get_column_letter(column_index)
                        )
                        fields.append(f"{header}: {value}")

                    content = (
                        f"工作表：{worksheet.title}\n"
                        f"第{row_number}行\n"
                        + "\n".join(fields)
                    )
                    block_type = "spreadsheet_row"

                documents.append(Document(
                    page_content=content,
                    metadata={
                        "source": filepath,
                        "block_type": block_type,
                        "sheet_name": worksheet.title,
                        "sheet_index": sheet_index,
                        "row_number": row_number,
                        "column_count": len(values),
                    },
                ))

        return documents
    except Exception as e:
        logger.error(f"解析XLSX文件 {filepath} 失败：{str(e)}", exc_info=True)
        raise
    finally:
        if workbook is not None:
            workbook.close()


def _format_excel_value(value) -> str:
    """Convert an Excel cell value into stable, searchable text."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, str):
        return " ".join(value.replace("\r\n", "\n").replace("\r", "\n").split())
    return str(value)


def _looks_like_excel_header(values) -> bool:
    """Return whether a row is suitable for use as a table header."""
    non_empty = [value.strip() for value in values if isinstance(value, str) and value.strip()]
    populated_count = sum(value is not None and str(value).strip() != "" for value in values)
    return (
        len(non_empty) >= 2
        and len(non_empty) == populated_count
        and len(non_empty) == len(set(non_empty))
        and all(not value.startswith("=") for value in non_empty)
    )


def docx_loader(filepath: str) -> list[Document]:
    """
    Word 文档加载器：解析 .docx 文件，提取结构化内容（标题、段落、表格）
    :param filepath: docx 文件的路径
    :return: Document 列表，每个 Document 代表一个内容块（标题/段落/表格）
    """
    try:
        doc = DocxDocument(filepath)
        blocks = []

        # 维护当前章节路径
        section_stack = []  # [(level, title), ...]
        paragraph_index = 0
        table_index = 0

        # 按原始顺序遍历文档内容
        for element in doc.element.body:
            if isinstance(element, CT_P):
                # 段落节点
                paragraph = Paragraph(element, doc)
                text = _extract_paragraph_text(paragraph)

                if not text.strip():
                    continue

                # 识别是否为标题
                heading_info = _identify_heading(paragraph, text)

                if heading_info:
                    # 标题块
                    level = heading_info['level']

                    # 更新章节栈
                    section_stack = _update_section_stack(section_stack, level, text)
                    section_path = _build_section_path(section_stack)

                    blocks.append(Document(
                        page_content=text,
                        metadata={
                            'block_type': 'heading',
                            'heading_level': level,
                            'is_heading': True,
                            'style_name': heading_info.get('style_name', ''),
                            'section_path': section_path,
                            'source': filepath
                        }
                    ))
                else:
                    # 段落块
                    section_path = _build_section_path(section_stack)

                    blocks.append(Document(
                        page_content=text,
                        metadata={
                            'block_type': 'paragraph',
                            'paragraph_index': paragraph_index,
                            'is_heading': False,
                            'section_path': section_path,
                            'style_name': paragraph.style.name if paragraph.style else '',
                            'source': filepath
                        }
                    ))
                    paragraph_index += 1

            elif isinstance(element, CT_Tbl):
                # 表格节点
                table = Table(element, doc)
                table_text = _extract_table_text(table)

                if not table_text.strip():
                    continue

                section_path = _build_section_path(section_stack)
                row_count = len(table.rows)
                column_count = len(table.columns) if table.rows else 0

                blocks.append(Document(
                    page_content=table_text,
                    metadata={
                        'block_type': 'table',
                        'table_index': table_index,
                        'row_count': row_count,
                        'column_count': column_count,
                        'section_path': section_path,
                        'source': filepath
                    }
                ))
                table_index += 1

        return blocks

    except Exception as e:
        logger.error(f"解析DOCX文件 {filepath} 失败：{str(e)}", exc_info=True)
        raise


def _extract_paragraph_text(paragraph: Paragraph) -> str:
    """提取段落的完整文本（合并所有 run）"""
    text = paragraph.text
    # 统一换行符
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    # 规范化空白，但保留型号中的连字符和空格
    text = ' '.join(text.split())
    return text.strip()


def _identify_heading(paragraph: Paragraph, text: str) -> Optional[dict]:
    """
    识别段落是否为标题
    :return: 标题信息 {'level': int, 'style_name': str} 或 None
    """
    # 优先使用样式识别
    if paragraph.style and paragraph.style.name:
        style_name = paragraph.style.name

        # 英文样式：Heading 1, Heading 2, ...
        if style_name.startswith('Heading'):
            match = re.match(r'Heading\s*(\d+)', style_name)
            if match:
                level = int(match.group(1))
                return {'level': level, 'style_name': style_name}

        # 中文样式：标题 1, 标题 2, ...
        if '标题' in style_name:
            match = re.search(r'标题\s*(\d+)', style_name)
            if match:
                level = int(match.group(1))
                return {'level': level, 'style_name': style_name}

    # 样式识别失败，使用保守的文本规则
    # 避免误判普通正文

    # Markdown 风格：# 标题
    if re.match(r'^#{1,6}\s+.+', text):
        level = len(re.match(r'^(#+)', text).group(1))
        return {'level': level, 'style_name': 'markdown'}

    # 数字编号：1. 产品介绍, 1.1 使用方法
    if re.match(r'^\d+(\.\d+)*\.?\s+.+', text):
        dots = text.split()[0].count('.')
        level = dots + 1 if dots > 0 else 1
        # 只有长度适中的才认为是标题（避免误判正文）
        if 2 <= len(text) <= 50:
            return {'level': level, 'style_name': 'numbered'}

    # 中文章节：第一章, 一、
    if re.match(r'^第[一二三四五六七八九十百]+章.+', text) or \
       re.match(r'^[一二三四五六七八九十]+、.+', text):
        if len(text) <= 50:
            return {'level': 1, 'style_name': 'chinese_chapter'}

    return None


def _update_section_stack(stack: list, new_level: int, title: str) -> list:
    """更新章节栈，保持层级关系"""
    # 移除所有大于等于当前级别的标题
    while stack and stack[-1][0] >= new_level:
        stack.pop()

    # 添加当前标题
    stack.append((new_level, title))
    return stack


def _build_section_path(stack: list) -> str:
    """构建章节路径字符串"""
    if not stack:
        return ''
    return ' > '.join([title for _, title in stack])


def _extract_table_text(table: Table) -> str:
    """
    提取表格内容，保留行列关系
    使用简单稳定的文本格式，不恢复复杂的合并单元格
    """
    if not table.rows:
        return ''

    lines = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            # 提取单元格文本
            cell_text = cell.text.strip()
            # 统一换行符，单元格内换行用空格代替
            cell_text = ' '.join(cell_text.split())
            cells.append(cell_text)

        # 使用 | 分隔单元格
        lines.append(' | '.join(cells))

    return '\n'.join(lines)
